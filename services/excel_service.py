"""Safe Excel import/export helpers used by customer and quotation pages."""

from __future__ import annotations

import json
import zipfile
from io import BytesIO
from typing import Any, Iterable

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from utils.validation import nonnegative_decimal, normalize_date

FORMULA_PREFIXES = ("=", "+", "-", "@")


def _safe_cell(value: Any) -> Any:
    """Prevent exported user text from being interpreted as an Excel formula."""
    if isinstance(value, (dict, list, tuple, set)):
        value = json.dumps(value, ensure_ascii=False, default=str)
    if isinstance(value, str) and value.startswith(FORMULA_PREFIXES):
        return "'" + value
    return value


def _heading(name: str) -> str:
    return name.replace("_", " ").title()


def _style_header(sheet) -> None:
    fill = PatternFill("solid", fgColor="0F6B78")
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        letter = column[0].column_letter
        longest = max((len(str(cell.value or "")) for cell in column), default=10)
        sheet.column_dimensions[letter].width = min(max(longest + 2, 12), 38)


def customers_to_excel(records: Iterable[dict[str, Any]]) -> bytes:
    """Export customer dictionaries to a readable XLSX workbook."""
    rows = list(records)
    preferred = [
        "company_name",
        "contact_name",
        "job_title",
        "country",
        "website",
        "email",
        "phone",
        "lead_source",
        "product_interest",
        "import_frequency",
        "estimated_purchase_volume",
        "last_contact_date",
        "next_follow_up_date",
        "current_stage",
        "lead_grade",
        "lead_score",
        "notes",
    ]
    other = sorted({key for row in rows for key in row} - set(preferred) - {"id"})
    columns = [column for column in preferred if any(column in row for row in rows)] + other
    if not columns:
        columns = preferred

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Customers"
    sheet.append([_heading(column) for column in columns])
    for row in rows:
        sheet.append([_safe_cell(row.get(column)) for column in columns])
    _style_header(sheet)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def import_customer_file(payload: bytes, filename: str) -> list[dict[str, Any]]:
    """Read a small CSV/XLSX upload and normalize column names."""
    if not payload:
        raise ValueError("The uploaded file is empty.")
    if len(payload) > 5 * 1024 * 1024:
        raise ValueError("The uploaded file exceeds the 5 MB limit.")
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    buffer = BytesIO(payload)
    if suffix == "csv":
        # Contact identifiers are text: inference drops leading zeroes in phones.
        # Keep explicit "NaN" input distinct from an empty optional value.
        frame = pd.read_csv(buffer, dtype=str, keep_default_na=False)
    elif suffix in {"xlsx", "xlsm"}:
        if not zipfile.is_zipfile(buffer):
            raise ValueError("The Excel file signature is invalid.")
        buffer.seek(0)
        with zipfile.ZipFile(buffer) as archive:
            members = archive.infolist()
            if len(members) > 200:
                raise ValueError("The Excel file contains too many internal files.")
            if sum(member.file_size for member in members) > 25 * 1024 * 1024:
                raise ValueError("The expanded Excel file exceeds the 25 MB safety limit.")
        buffer.seek(0)
        frame = pd.read_excel(buffer, engine="openpyxl")
    else:
        raise ValueError("Only CSV and XLSX files are supported.")
    if frame.empty:
        raise ValueError("The uploaded file has no customer rows.")
    if len(frame) > 5_000 or len(frame.columns) > 50:
        raise ValueError("Customer imports are limited to 5,000 rows and 50 columns.")
    frame.columns = [
        str(column).strip().lower().replace(" / ", "_").replace(" ", "_").replace("-", "_")
        for column in frame.columns
    ]
    if frame.columns.duplicated().any():
        raise ValueError("Column names must be unique after normalization.")
    if "company_name" not in frame.columns:
        raise ValueError("The file must contain a Company Name column.")
    # Float columns otherwise retain NaN when replacing missing cells with None.
    frame = frame.astype(object).where(pd.notna(frame), None)
    normalized_rows: list[dict[str, Any]] = []
    for row_number, row in enumerate(frame.to_dict(orient="records"), start=2):
        normalized = {
            key: value.strip() if isinstance(value, str) else value
            for key, value in row.items()
        }
        volume = normalized.get("estimated_purchase_volume")
        if volume not in (None, ""):
            numeric_text = volume.replace(",", "") if isinstance(volume, str) else volume
            try:
                normalized["estimated_purchase_volume"] = float(
                    nonnegative_decimal(numeric_text, "Estimated Purchase Volume")
                )
            except ValueError as exc:
                raise ValueError(f"Row {row_number}: {exc}") from exc
        else:
            normalized["estimated_purchase_volume"] = 0.0
        for field, label in (
            ("last_contact_date", "Last Contact Date"),
            ("next_follow_up_date", "Next Follow-up Date"),
        ):
            try:
                normalized[field] = normalize_date(normalized.get(field))
            except (TypeError, ValueError) as exc:
                raise ValueError(
                    f"Row {row_number}: {label} must be a valid ISO date (YYYY-MM-DD)."
                ) from exc
        normalized_rows.append(normalized)
    return normalized_rows


def quotation_to_excel(quote: dict[str, Any], metadata: dict[str, Any]) -> bytes:
    """Create a transparent quotation workbook without hidden formulas."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Quotation"
    sheet.append(["Quotation Summary", "Value"])

    method = quote["pricing_method"]
    method_label = "Gross Margin" if method == "gross_margin" else "Markup"
    formula = (
        "Selling price = cost / (1 - gross margin rate)"
        if method == "gross_margin"
        else "Selling price = cost * (1 + markup rate)"
    )
    summary = [
        ("Product", metadata.get("product_name", "")),
        ("Pricing Method", method_label),
        ("Formula", formula),
        ("Rate", f"{float(quote.get('pricing_rate', quote.get('rate', 0))) * 100:.2f}%"),
        ("Quantity", quote["quantity"]),
        ("Exchange Rate", f"1 USD = {quote.get('exchange_rate_cny_per_usd', quote.get('exchange_rate'))} CNY"),
        ("Valid Until", metadata.get("valid_until", "")),
        ("MOQ", metadata.get("moq", "")),
        ("Lead Time", metadata.get("lead_time", "")),
        ("Payment Terms", metadata.get("payment_terms", "")),
        ("Notes", metadata.get("notes", "")),
    ]
    for label, value in summary:
        sheet.append([label, _safe_cell(value)])

    sheet.append([])
    sheet.append(["Incoterm", "Total USD", "Unit USD", "Gross Profit USD", "Gross Margin %"])
    for term, result in quote["terms"].items():
        sheet.append(
            [
                term,
                float(result["total_usd"]),
                float(result["unit_usd"]),
                float(result["gross_profit_usd"]),
                float(result.get("gross_margin_percent", float(result.get("gross_margin", 0)) * 100)),
            ]
        )

    header_fill = PatternFill("solid", fgColor="0F6B78")
    for row_number in (1, 14):
        for cell in sheet[row_number]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
    sheet.column_dimensions["A"].width = 26
    sheet.column_dimensions["B"].width = 52
    sheet.column_dimensions["C"].width = 18
    sheet.column_dimensions["D"].width = 20
    sheet.column_dimensions["E"].width = 18
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
