from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook, load_workbook

import pytest

from services.excel_service import customers_to_excel, import_customer_file, quotation_to_excel


def test_customer_excel_export_can_be_reopened() -> None:
    payload = customers_to_excel(
        [
            {
                "company_name": "Fictional Atlas Paper Studio",
                "contact_name": "Demo Contact",
                "email": "demo.contact@example.com",
                "notes": "=HYPERLINK(\"bad\",\"bad\")",
                "score_breakdown": {"company_authenticity": 18},
            }
        ]
    )

    workbook = load_workbook(BytesIO(payload))
    sheet = workbook["Customers"]
    assert sheet.max_row == 2
    headers = [cell.value for cell in sheet[1]]
    notes_col = headers.index("Notes") + 1
    assert str(sheet.cell(row=2, column=notes_col).value).startswith("'")
    breakdown_col = headers.index("Score Breakdown") + 1
    assert "company_authenticity" in sheet.cell(row=2, column=breakdown_col).value


def test_quotation_excel_contains_formula_explanation() -> None:
    quote = {
        "pricing_method": "gross_margin",
        "rate": "0.20",
        "exchange_rate": "7.20",
        "quantity": 1000,
        "terms": {
            "FOB": {
                "total_usd": "1250.00",
                "unit_usd": "1.25",
                "gross_profit_usd": "250.00",
                "gross_margin_percent": "20.00",
                "components": {"product_cost": "6000.00"},
            }
        },
    }

    payload = quotation_to_excel(
        quote,
        {
            "product_name": "Fictional custom notebook",
            "valid_until": "2026-08-15",
            "moq": 1000,
            "lead_time": "25 days",
            "payment_terms": "30% deposit, balance before shipment",
        },
    )
    workbook = load_workbook(BytesIO(payload), data_only=False)
    assert "Quotation" in workbook.sheetnames
    values = [cell.value for row in workbook["Quotation"].iter_rows() for cell in row]
    assert "Gross Margin" in values
    assert "Selling price = cost / (1 - gross margin rate)" in values


def test_invalid_excel_signature_is_rejected() -> None:
    with pytest.raises(ValueError, match="signature"):
        import_customer_file(b"not-a-zip-workbook", "customers.xlsx")


@pytest.mark.parametrize(
    ("column", "value", "message"),
    [
        ("Estimated Purchase Volume", "abc", "must be a number"),
        ("Estimated Purchase Volume", "NaN", "finite number"),
        ("Last Contact Date", "not-a-date", "valid ISO date"),
        ("Next Follow-up Date", "also-bad", "valid ISO date"),
    ],
)
def test_customer_import_rejects_invalid_business_values(
    column: str, value: str, message: str
) -> None:
    csv_payload = f"Company Name,{column}\nBad CSV Buyer,{value}\n".encode()

    with pytest.raises(ValueError, match=message):
        import_customer_file(csv_payload, "customers.csv")


def test_customer_import_normalizes_volume_and_dates() -> None:
    payload = (
        b"Company Name,Estimated Purchase Volume,Last Contact Date,Next Follow-up Date\n"
        b"Fictional CSV Buyer,12000,2026-07-01,2026-07-20\n"
    )

    row = import_customer_file(payload, "customers.csv")[0]

    assert row["estimated_purchase_volume"] == 12000.0
    assert row["last_contact_date"] == "2026-07-01"
    assert row["next_follow_up_date"] == "2026-07-20"


def test_csv_import_preserves_contact_identifiers_and_optional_blank_volume() -> None:
    payload = (
        b"Company Name,Phone,Estimated Purchase Volume\n"
        b"Fictional Buyer One,001234567,120\n"
        b"Fictional Buyer Two,009876543,\n"
    )

    rows = import_customer_file(payload, "customers.csv")

    assert [row["phone"] for row in rows] == ["001234567", "009876543"]
    assert [row["estimated_purchase_volume"] for row in rows] == [120.0, 0.0]


def test_excel_import_accepts_blank_optional_numeric_cells() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Company Name", "Estimated Purchase Volume"])
    sheet.append(["Fictional Buyer One", 120])
    sheet.append(["Fictional Buyer Two", None])
    buffer = BytesIO()
    workbook.save(buffer)

    rows = import_customer_file(buffer.getvalue(), "customers.xlsx")

    assert [row["estimated_purchase_volume"] for row in rows] == [120.0, 0.0]


def test_import_rejects_columns_that_would_overwrite_customer_values() -> None:
    payload = b"Company Name,Company-Name\nFictional Buyer One,Fictional Buyer Two\n"

    with pytest.raises(ValueError, match="Column names must be unique"):
        import_customer_file(payload, "customers.csv")
